import json
import logging
import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook

from .models import CatalogueDataset, CatalogueEntry, Note, Programme, TrackerDataset

log = logging.getLogger(__name__)

DATA_DIR = Path(os.environ.get("DATA_DIR") or Path(__file__).resolve().parents[2])
FEES_FILE = Path(__file__).resolve().parents[1] / "data" / "fees.json"


class DataError(Exception):
    pass


@dataclass(frozen=True)
class Source:
    id: str
    label: str
    kind: Literal["tracker", "catalogue"]
    file: str
    sheet: str
    readme_sheet: str | None = None
    window_sheet: str | None = None


SOURCES = (
    Source("direct", "Direct CS / AI", "tracker", "Direct_CS.xlsx", "All Direct CS-AI", readme_sheet="Read Me"),
    Source(
        "adjacent",
        "CS-Adjacent",
        "tracker",
        "InDirect_CS.xlsx",
        "All Remaining Adjacent",
        readme_sheet="Read Me",
        window_sheet="Open Oct28-Dec31",
    ),
    Source("catalogue", "Catalogue", "catalogue", "All_Scholarships.xlsx", "All EMJM"),
)

TRACKER_COLUMNS = {
    "Name": "name",
    "Countries / Mobility": "countries",
    "Start date": "opens",
    "Start date (2027 intake applications)": "opens",
    "End date": "deadline",
    "End date (scholarship deadline)": "deadline",
    "Priority Level (1–3)": "priority",
    "Admission Type": "admission_type",
    "Admission Type (Both/Single)": "admission_type",
    "Stipend / Funding": "stipend",
    "IELTS Bands": "ielts",
    "MOI Accepted?": "moi",
    "GRE": "gre",
    "Pakistan Eligible": "pakistan_eligible",
    "University Course / Admission Link": "admission_url",
    "Erasmus Apply Link": "apply_url",
    "Suggestions / Ideal Candidates": "suggestions",
    "2027 Call Status": "call_status",
    "Official Verification Source": "source_url",
}

CATALOGUE_COLUMNS = {
    "Programme title": "title",
    "Acronym": "acronym",
    "Subject area (title-based)": "subject",
    "CS / AI relevance": "relevance",
    "Funding cycle / source status": "cycle",
    "2027 application status": "application_status",
    "2027 scholarship status": "scholarship_status",
    "Application deadline": "deadline",
    "Duration / ECTS": "duration",
    "Consortium / study countries": "countries",
    "Language of instruction": "language",
    "Academic eligibility / prerequisites": "eligibility",
    "Funding / coverage detail": "funding",
    "Official programme website": "website_url",
    "EU project record": "project_url",
    "EU catalogue / selection source": "catalogue_url",
    "Checked on": "checked_on",
    "Fit / verification notes": "notes",
}

TRACKER_TEXT = (
    "countries", "opens", "deadline", "admission_type", "stipend", "ielts", "moi", "gre",
    "pakistan_eligible", "admission_url", "apply_url", "call_status", "source_url",
)  # fmt: skip
CATALOGUE_TEXT = (
    "title", "subject", "relevance", "cycle", "application_status", "scholarship_status", "deadline",
    "duration", "countries", "language", "eligibility", "funding", "website_url", "project_url",
    "catalogue_url", "checked_on",
)  # fmt: skip

ISO_DAY = re.compile(r"^\d{4}-\d{2}-\d{2}")


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _day(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = ISO_DAY.match(_text(value))
    if not match:
        return None
    try:
        return date.fromisoformat(match.group())
    except ValueError:
        return None


def _priority(value: Any) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _split(value: Any, sep: str) -> list[str]:
    return [part.strip() for part in _text(value).split(sep) if part.strip()]


def _split_name(name: str) -> tuple[str | None, str]:
    acronym, sep, title = name.partition(" — ")
    return (acronym.strip(), title.strip()) if sep else (None, name)


def _mapped_rows(wb, sheet: str, columns: dict[str, str]) -> list[dict[str, Any]]:
    if sheet not in wb.sheetnames:
        raise DataError(f"Sheet '{sheet}' not found (available: {', '.join(wb.sheetnames)})")
    rows = wb[sheet].iter_rows(values_only=True)
    header = [_text(h) for h in next(rows, ())]
    unknown = [h for h in header if h and h not in columns]
    if unknown:
        log.warning("Sheet %r has unmapped columns: %s", sheet, ", ".join(unknown))
    fields = [columns.get(h) for h in header]
    return [{f: v for f, v in zip(fields, row) if f} for row in rows if any(_text(v) for v in row)]


def _notes(wb, sheet: str) -> list[Note]:
    if sheet not in wb.sheetnames:
        raise DataError(f"Sheet '{sheet}' not found (available: {', '.join(wb.sheetnames)})")
    return [
        Note(label=_text(label), text=_text(text))
        for label, text, *_ in wb[sheet].iter_rows(values_only=True)
        if _text(label) and _text(text)
    ]


def _load_fees() -> dict[str, dict[str, str]]:
    try:
        return json.loads(FEES_FILE.read_text(encoding="utf-8"))
    except FileNotFoundError:
        log.warning("No fee data at %s", FEES_FILE)
        return {}


def _programme(source_id: str, index: int, row: dict[str, Any], window: set[str], fees: dict[str, dict[str, str]]) -> Programme:
    name = _text(row.get("name"))
    acronym, title = _split_name(name)
    return Programme(
        id=f"{source_id}-{index}",
        acronym=acronym,
        title=title,
        priority=_priority(row.get("priority")),
        opens_on=_day(row.get("opens")),
        deadline_on=_day(row.get("deadline")),
        suggestions=_split(row.get("suggestions"), "•"),
        in_window=name in window,
        **fees.get(f"{source_id}:{acronym or title}", {}),
        **{field: _text(row.get(field)) for field in TRACKER_TEXT},
    )


def _attach_alerts(programmes: list[Programme], notes: list[Note]) -> list[Note]:
    by_acronym = {p.acronym.lower(): p for p in programmes if p.acronym}
    general = []
    for note in notes:
        target = by_acronym.get(note.label.split()[0].lower())
        if target:
            target.alerts.append(note)
        else:
            general.append(note)
    return general


def _build_tracker(source: Source, wb) -> TrackerDataset:
    rows = _mapped_rows(wb, source.sheet, TRACKER_COLUMNS)
    window = set()
    if source.window_sheet:
        window = {_text(r.get("name")) for r in _mapped_rows(wb, source.window_sheet, TRACKER_COLUMNS)}
    fees = _load_fees()
    programmes = [_programme(source.id, i, row, window, fees) for i, row in enumerate(rows)]
    matched = {f"{source.id}:{p.acronym or p.title}" for p in programmes}
    stale = [k for k in fees if k.startswith(f"{source.id}:") and k not in matched]
    if stale:
        log.warning("fees.json has entries matching no %s programme: %s", source.id, ", ".join(stale))
    notes = _notes(wb, source.readme_sheet) if source.readme_sheet else []
    return TrackerDataset(id=source.id, label=source.label, notes=_attach_alerts(programmes, notes), items=programmes)


def _build_catalogue(source: Source, wb) -> CatalogueDataset:
    entries = []
    for i, row in enumerate(_mapped_rows(wb, source.sheet, CATALOGUE_COLUMNS)):
        entries.append(
            CatalogueEntry(
                id=f"{source.id}-{i}",
                acronym=_text(row.get("acronym")) or None,
                notes=_split(row.get("notes"), " | "),
                **{field: _text(row.get(field)) for field in CATALOGUE_TEXT},
            )
        )
    return CatalogueDataset(id=source.id, label=source.label, notes=[], items=entries)


def _build(source: Source) -> TrackerDataset | CatalogueDataset:
    wb = load_workbook(DATA_DIR / source.file, read_only=True, data_only=True)
    try:
        return _build_tracker(source, wb) if source.kind == "tracker" else _build_catalogue(source, wb)
    finally:
        wb.close()


_cache: dict[str, tuple[tuple[int, int], TrackerDataset | CatalogueDataset]] = {}


def find_source(dataset_id: str) -> Source | None:
    return next((s for s in SOURCES if s.id == dataset_id), None)


def get_dataset(source: Source) -> TrackerDataset | CatalogueDataset:
    path = DATA_DIR / source.file
    try:
        stamp = (path.stat().st_mtime_ns, FEES_FILE.stat().st_mtime_ns if FEES_FILE.exists() else 0)
    except FileNotFoundError:
        raise DataError(f"File not found: {path}") from None
    cached = _cache.get(source.id)
    if cached and cached[0] == stamp:
        return cached[1]
    dataset = _build(source)
    _cache[source.id] = (stamp, dataset)
    return dataset
